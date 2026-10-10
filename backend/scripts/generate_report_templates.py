"""Generate the seven versioned, empty Excel templates from fixed definitions."""

import io
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from app.services.report_service import build_xlsx_report


KINDS = ("ATTENDANCE", "LEAVE", "ATTENDANCE_FIX", "APPROVAL_QUEUE", "HEADCOUNT", "MY_PAYSLIP", "PAYROLL_SUMMARY")


def main():
    directory = Path(__file__).resolve().parents[1] / "app" / "report_templates"
    directory.mkdir(parents=True, exist_ok=True)
    for kind in KINDS:
        raw = build_xlsx_report({"kind": kind, "start_date": date(2000, 1, 1), "end_date": date(2000, 1, 1), "scope": "SELF", "rows": [], "note": ""}, use_template=False)
        workbook = load_workbook(io.BytesIO(raw))
        for row in workbook["Thông tin"].iter_rows(min_row=2):
            row[1].value = None
        workbook["Tổng hợp"].delete_rows(2, workbook["Tổng hợp"].max_row)
        path = directory / f"{kind.lower()}_v1.xlsx"
        workbook.save(path)
        print(path.name)


if __name__ == "__main__":
    main()
