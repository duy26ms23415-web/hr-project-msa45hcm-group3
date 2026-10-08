from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException
from sqlalchemy.dialects import postgresql

from app.services.report_service import ReportService, build_xlsx_report
from app.services.report_storage import delete_report, resolve_report_path, store_report
from app.api.v1.endpoints.reports import _authorized_run


def user(role: str):
    return SimpleNamespace(
        employee_id=7,
        login_email="manager@example.test",
        role_assignments=[SimpleNamespace(role=SimpleNamespace(role_code=role))],
    )


def empty_rows():
    return SimpleNamespace(mappings=lambda: SimpleNamespace(all=lambda: []))


@pytest.mark.asyncio
@pytest.mark.parametrize("kind,expected_calls", [
    ("ATTENDANCE_FIX", 1),
    ("APPROVAL_QUEUE", 2),
    ("HEADCOUNT", 1),
])
async def test_additional_reports_filter_manager_scope_before_return(kind, expected_calls):
    db = AsyncMock()
    db.execute.return_value = empty_rows()
    report = await ReportService.generate(db, user("MANAGER"), kind, date(2026, 10, 1), date(2026, 10, 31), None)
    assert report["kind"] == kind and report["scope"] == "DIRECT_REPORTS" and report["rows"] == []
    assert db.execute.await_count == expected_calls
    for call in db.execute.await_args_list:
        sql = str(call.args[0].compile(dialect=postgresql.dialect()))
        assert "manager_employee_id" in sql


@pytest.mark.asyncio
async def test_employee_cannot_create_headcount_report():
    db = AsyncMock()
    with pytest.raises(HTTPException) as denied:
        await ReportService.generate(db, user("EMPLOYEE"), "HEADCOUNT", date(2026, 10, 1), date(2026, 10, 31), None)
    assert denied.value.status_code == 403
    db.execute.assert_not_called()


def test_xlsx_template_has_expected_sheets_and_neutralizes_formulas():
    from io import BytesIO
    from openpyxl import load_workbook

    report = {
        "kind": "ATTENDANCE",
        "start_date": date(2026, 10, 1),
        "end_date": date(2026, 10, 31),
        "scope": "SELF",
        "note": "Recorded rows only",
        "rows": [{
            "employee_code": "=1+1", "full_name": "@SUM(A1)", "department_id": 3,
            "recorded_days": 4, "present_days": 3, "incomplete_days": 1,
            "absent_days": 0, "worked_minutes": 123,
        }],
    }
    workbook = load_workbook(BytesIO(build_xlsx_report(report)), data_only=False)
    assert workbook.sheetnames == ["Thông tin", "Dữ liệu", "Tổng hợp"]
    data = workbook["Dữ liệu"]
    assert data["A2"].value == "'=1+1" and data["B2"].value == "'@SUM(A1)"
    assert data["H2"].value == 123 and data["H2"].data_type == "n" and data["A2"].data_type != "f"
    report["kind"] = "LEAVE"
    report["rows"] = [{"employee_code": "E-7", "full_name": "Example", "department_id": 3, "approved_days": Decimal("0.50"), "pending_requests": 1, "rejected_requests": 0, "remaining_annual_days": Decimal("2.25")}]
    leave_data = load_workbook(BytesIO(build_xlsx_report(report)), data_only=False)["Dữ liệu"]
    assert leave_data["D2"].value == 0.5 and leave_data["D2"].data_type == "n"
    assert leave_data["G2"].value == 2.25 and leave_data["G2"].data_type == "n"


def test_private_report_artifact_storage_rejects_untrusted_paths(tmp_path, monkeypatch):
    from app.core.config import settings
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", tmp_path)
    key, digest = store_report(b"workbook-bytes", "xlsx")
    path = resolve_report_path(key)
    assert path.read_bytes() == b"workbook-bytes" and len(digest) == 64
    with pytest.raises(FileNotFoundError):
        resolve_report_path("..\\outside.xlsx")
    delete_report(key)
    assert not path.exists()


def test_orphan_scan_ignores_referenced_recent_and_nonserver_files(tmp_path, monkeypatch):
    import os
    from app.core.config import settings
    from app.services.report_storage import orphan_report_keys
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", tmp_path)
    old = (datetime.now(timezone.utc) - timedelta(hours=3)).timestamp()
    referenced, orphan, recent = "a" * 32 + ".xlsx", "b" * 32 + ".json", "c" * 32 + ".xlsx"
    for key in (referenced, orphan, recent, "unknown.xlsx"):
        (tmp_path / key).write_bytes(b"artifact")
        if key != recent:
            os.utime(tmp_path / key, (old, old))
    assert orphan_report_keys({referenced}) == [orphan]
    assert (tmp_path / orphan).exists()  # Scan itself never deletes.


@pytest.mark.asyncio
async def test_report_run_read_rechecks_owner_and_current_employee_scope(monkeypatch):
    run = SimpleNamespace(
        run_id="12345678123456781234567812345678", owner_user_account_id=9,
        status="READY", expires_at=datetime.now(timezone.utc) + timedelta(minutes=10),
        kind="ATTENDANCE", scope="SELF", filters={"start_date": "2026-10-01", "end_date": "2026-10-31"}, scope_snapshot=[7],
    )
    db = AsyncMock()
    db.scalar.return_value = run
    monkeypatch.setattr(ReportService, "visible_employee_ids", AsyncMock(return_value=("SELF", [7])))
    owner = SimpleNamespace(user_account_id=9)
    assert await _authorized_run(db, owner, run.run_id) is run
    sql = str(db.scalar.call_args.args[0].compile(dialect=postgresql.dialect()))
    assert "owner_user_account_id" in sql

    monkeypatch.setattr(ReportService, "visible_employee_ids", AsyncMock(return_value=("SELF", [8])))
    with pytest.raises(HTTPException) as denied:
        await _authorized_run(db, owner, run.run_id)
    assert denied.value.status_code == 404


@pytest.mark.asyncio
async def test_owner_report_expiry_returns_expired_code_before_data_lookup(monkeypatch):
    run = SimpleNamespace(status="READY", expires_at=datetime.now(timezone.utc) - timedelta(seconds=1), storage_key=None, snapshot_storage_key=None)
    db = AsyncMock()
    db.scalar.return_value = run
    scope = AsyncMock()
    monkeypatch.setattr(ReportService, "visible_employee_ids", scope)
    with pytest.raises(HTTPException) as expired:
        await _authorized_run(db, SimpleNamespace(user_account_id=9), "a" * 32)
    assert expired.value.status_code == 410 and expired.value.detail == "REPORT_EXPIRED"
    assert run.status == "EXPIRED"
    scope.assert_not_called()


@pytest.mark.asyncio
async def test_failed_run_is_not_reclassified_by_preview_after_expiry():
    run = SimpleNamespace(status="FAILED", expires_at=datetime.now(timezone.utc) - timedelta(days=1))
    db = AsyncMock()
    db.scalar.return_value = run
    with pytest.raises(HTTPException) as failed:
        await _authorized_run(db, SimpleNamespace(user_account_id=9), "a" * 32)
    assert failed.value.status_code == 404 and run.status == "FAILED"
    db.commit.assert_not_called()


@pytest.mark.asyncio
async def test_missing_employee_and_invalid_queue_scope_fail_before_query():
    db = AsyncMock()
    hr = user("HR")
    hr.employee_id = None
    for kind, actor, scope in [("MY_PAYSLIP", hr, "SELF"), ("APPROVAL_QUEUE", user("MANAGER"), "COMPANY")]:
        with pytest.raises(HTTPException) as rejected:
            await ReportService.generate(db, actor, kind, date(2026, 10, 1), date(2026, 10, 31), None, scope)
        assert rejected.value.status_code == 403
    db.execute.assert_not_called()
    db.scalar.assert_not_called()


@pytest.mark.asyncio
async def test_payroll_report_denies_manager_summary_before_query_and_unreleased_period_before_lines():
    db = AsyncMock()
    with pytest.raises(HTTPException) as denied:
        await ReportService.generate(db, user("MANAGER"), "PAYROLL_SUMMARY", date(2026, 10, 1), date(2026, 10, 31), None)
    assert denied.value.status_code == 403
    db.execute.assert_not_called()
    db.scalar.assert_not_called()
    db.scalar.return_value = SimpleNamespace(status="CALCULATED")
    with pytest.raises(HTTPException) as hidden:
        await ReportService.generate(db, user("EMPLOYEE"), "MY_PAYSLIP", date(2026, 10, 1), date(2026, 10, 31), None)
    assert hidden.value.status_code == 404
    db.execute.assert_not_called()


@pytest.mark.asyncio
@pytest.mark.parametrize("status", ["APPROVED", "CLOSED"])
async def test_released_payslip_filters_authenticated_employee_and_preserves_decimal(status):
    db = AsyncMock()
    db.scalar.return_value = SimpleNamespace(status=status, payroll_period_id=1, period_year=2026, period_month=10)
    fields = ("base_salary", "standard_work_days", "actual_work_days", "allowance_amount", "overtime_amount", "deduction_amount", "gross_salary", "insurance_deduction", "taxable_income", "personal_income_tax", "other_deductions", "net_salary")
    line = SimpleNamespace(employee_id=7, currency_code="VND", **{field: Decimal("1234567890123456.78") for field in fields})
    db.execute.return_value = SimpleNamespace(first=lambda: (line, "E7", "Employee Seven", 1))
    report = await ReportService.generate(db, user("EMPLOYEE"), "MY_PAYSLIP", date(2026, 10, 1), date(2026, 10, 31), None)
    sql = str(db.execute.call_args.args[0].compile(dialect=postgresql.dialect(), compile_kwargs={"literal_binds": True}))
    assert "hr_payroll_lines.employee_id = 7" in sql
    assert report["rows"][0]["net_salary"] == "1234567890123456.78"
    from io import BytesIO
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(build_xlsx_report(report)), data_only=False)
    assert workbook["Dữ liệu"]["Q2"].value == "1234567890123456.78"
    assert workbook["Dữ liệu"]["Q2"].data_type == "s"


@pytest.mark.asyncio
async def test_payroll_summary_keeps_currencies_separate_in_rows_and_excel_totals():
    fields = ("base_salary", "allowance_amount", "overtime_amount", "deduction_amount", "gross_salary", "insurance_deduction", "taxable_income", "personal_income_tax", "other_deductions", "net_salary")
    lines = [(SimpleNamespace(currency_code=currency, **{field: Decimal(amount) for field in fields}), "E7", "Employee", 1, "Department") for currency, amount in [("VND", "1.25"), ("USD", "2.50"), ("VND", "3.50")]]
    db = AsyncMock()
    db.scalar.return_value = SimpleNamespace(status="CLOSED", payroll_period_id=1)
    db.execute.return_value = SimpleNamespace(all=lambda: lines)
    report = await ReportService.generate(db, user("HR"), "PAYROLL_SUMMARY", date(2026, 10, 1), date(2026, 10, 31), None)
    assert {row["currency_code"]: row["net_salary"] for row in report["rows"]} == {"VND": "4.75", "USD": "2.50"}
    from io import BytesIO
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(build_xlsx_report(report)), data_only=False)
    summary = dict(workbook["Tổng hợp"].values)
    assert summary["Lương thực nhận NET (VND)"] == 4.75
    assert summary["Lương thực nhận NET (USD)"] == 2.50
    assert "Lương thực nhận NET" not in summary


@pytest.mark.asyncio
async def test_unknown_legacy_currency_is_not_labeled_or_aggregated_as_vnd():
    db = AsyncMock()
    db.scalar.return_value = SimpleNamespace(status="APPROVED", payroll_period_id=1)
    db.execute.return_value = SimpleNamespace(first=lambda: (SimpleNamespace(currency_code=None), "E7", "Employee", 1))
    with pytest.raises(HTTPException) as unknown:
        await ReportService.generate(db, user("EMPLOYEE"), "MY_PAYSLIP", date(2026, 10, 1), date(2026, 10, 31), None)
    assert unknown.value.status_code == 409 and unknown.value.detail == "PAYROLL_CURRENCY_UNAVAILABLE"



def test_workbook_uses_print_layout_native_dates_and_readable_scope():
    from io import BytesIO
    from openpyxl import load_workbook
    report = {"kind": "ATTENDANCE_FIX", "start_date": date(2026, 10, 1),
              "end_date": date(2026, 10, 8), "scope": "DIRECT_REPORTS",
              "note": "Recorded data only", "rows": [{
                  "attendance_fix_id": 1, "employee_code": "E01", "full_name": "Example",
                  "department_id": 1, "work_date": date(2026, 10, 5),
                  "event_type": "CHECK_IN", "requested_at": datetime(2026, 10, 5, 1, tzinfo=timezone.utc),
                  "status": "APPROVED", "reason": "=malicious",
              }]}
    workbook = load_workbook(BytesIO(build_xlsx_report(report)))
    data = workbook["Dữ liệu"]
    assert data.freeze_panes == "C2"
    assert data.page_setup.orientation == "landscape" and data.page_setup.fitToWidth == 1
    assert data.print_title_rows == "$1:$1"
    assert data["E2"].number_format == "dd/mm/yyyy"
    assert data["G2"].value.hour == 8
    assert data["I2"].value == "'=malicious"
    assert "ReportData" in data.tables and data.tables["ReportData"].ref == "A1:I2"
    assert workbook["Thông tin"]["B5"].value == "Nhân viên trực tiếp"
