"""Report metadata derived from shipped templates and current actor capabilities."""

from functools import lru_cache
from pathlib import Path
from openpyxl import load_workbook

from app.services.ai_access import require_known_role, user_roles


REPORT_NAMES = {
    "ATTENDANCE": "Chấm công", "LEAVE": "Nghỉ phép", "ATTENDANCE_FIX": "Giải trình công",
    "APPROVAL_QUEUE": "Đơn chờ tôi duyệt", "HEADCOUNT": "Nhân sự hiện tại",
    "MY_PAYSLIP": "Phiếu lương của tôi", "PAYROLL_SUMMARY": "Tổng hợp lương đã phát hành",
}


@lru_cache(maxsize=7)
def template_columns(kind):
    if kind not in REPORT_NAMES:
        raise ValueError("REPORT_KIND_UNSUPPORTED")
    path = Path(__file__).resolve().parents[1] / "report_templates" / f"{kind.lower()}_v1.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return tuple(cell.value for cell in next(workbook["Dữ liệu"].iter_rows(min_row=1, max_row=1)))
    finally:
        workbook.close()


def report_catalog(user):
    roles = user_roles(user)
    require_known_role(roles)
    employee = user.employee_id is not None
    privileged = bool(roles & {"HR", "ADMIN"})
    reviewer = bool(roles & {"MANAGER", "HR", "ADMIN"})
    common_scopes = (["SELF"] if employee else []) + (["DIRECT_REPORTS"] if reviewer and employee else []) + (["COMPANY"] if privileged else [])
    items = []
    for kind, label in REPORT_NAMES.items():
        scopes = common_scopes
        if kind == "MY_PAYSLIP":
            scopes = ["SELF"] if employee else []
        elif kind == "PAYROLL_SUMMARY":
            scopes = ["COMPANY"] if privileged else []
        elif kind == "APPROVAL_QUEUE":
            scopes = ["DIRECT_REPORTS"] if reviewer and employee else []
        elif kind == "HEADCOUNT" and not reviewer:
            scopes = []
        if not scopes:
            continue
        payroll = kind in {"MY_PAYSLIP", "PAYROLL_SUMMARY"}
        items.append({
            "kind": kind, "label": label, "template_version": f"{kind.lower()}_v1",
            "columns": list(template_columns(kind)), "scopes": scopes,
            "default_scope": "COMPANY" if kind in {"HEADCOUNT", "PAYROLL_SUMMARY"} and privileged else "DIRECT_REPORTS" if kind in {"HEADCOUNT", "APPROVAL_QUEUE"} else "SELF" if employee else scopes[0],
            "filters": ["start_date", "end_date"] + ([] if kind == "MY_PAYSLIP" else ["department_id"]),
            "period_rule": "CALENDAR_MONTH" if payroll else "SAME_YEAR",
            "released_statuses": ["APPROVED", "CLOSED"] if payroll else [],
            "max_rows": 10000,
        })
    return {"catalog_version": "v1", "reports": items}
