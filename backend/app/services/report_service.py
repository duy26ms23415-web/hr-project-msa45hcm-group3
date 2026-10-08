import csv
import io
import re
from pathlib import Path
from datetime import date, datetime, timezone
from decimal import Decimal
from zoneinfo import ZoneInfo
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from sqlalchemy import select, func, case, and_
from fastapi import HTTPException
from app.models.organization import Employee, Department
from app.models.attendance import AttendanceDay, AttendanceFix
from app.models.leave import LeaveRequest, EmployeeLeaveBalance
from app.models.payroll import PayrollPeriod, PayrollLine
from app.services.ai_access import report_scope, user_roles

MAX_REPORT_ROWS = 10000


def check_report_size(rows):
    if len(rows) > MAX_REPORT_ROWS:
        raise HTTPException(status_code=422, detail="REPORT_TOO_LARGE")
    return rows


def employee_scope(user, scope: str | None = None):
    selected = report_scope(user, scope)
    if selected == "COMPANY":
        return None
    if selected == "DIRECT_REPORTS":
        return Employee.manager_employee_id == user.employee_id
    return Employee.employee_id == user.employee_id


def safe_csv(rows, fields):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(fields)
    for row in rows:
        values = []
        for field in fields:
            value = "" if row.get(field) is None else str(row[field])
            if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
                value = "'" + value
            values.append(value)
        writer.writerow(values)
    return "\ufeff" + stream.getvalue()


def safe_xlsx_text(value) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def xlsx_decimal(value):
    number = Decimal(str(value))
    exact = format(number, "f")
    # Excel guarantees only 15 significant decimal digits. Keep larger values
    # as exact text instead of allowing the writer/viewer to round them.
    digits = exact.replace(".", "").lstrip("-0")
    return exact if len(digits) > 15 else number


def xlsx_value(field: str, value):
    if value is None:
        return None
    if field in {"department_id", "attendance_fix_id", "request_id", "employee_count", "period_year", "period_month", "recorded_days", "present_days", "incomplete_days", "absent_days", "worked_minutes", "pending_requests", "rejected_requests"}:
        return int(value)
    if field in {"approved_days", "remaining_annual_days"}:
        return xlsx_decimal(value)
    if field in {
        "standard_work_days", "actual_work_days", "base_salary", "allowance_amount", "overtime_amount", "deduction_amount", "gross_salary",
        "insurance_deduction", "taxable_income", "personal_income_tax", "other_deductions", "net_salary",
    }:
        return xlsx_decimal(value)
    if field in {"work_date", "request_date"}:
        return value if isinstance(value, date) else date.fromisoformat(str(value))
    if field in {"requested_at", "created_at"}:
        stamp = value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
        # Persist UTC, present office time as an Excel datetime without tzinfo.
        if stamp.tzinfo is None:
            stamp = stamp.replace(tzinfo=timezone.utc)
        return stamp.astimezone(ZoneInfo("Asia/Ho_Chi_Minh")).replace(tzinfo=None)
    vocabulary = {"CHECK_IN": "Vào làm", "CHECK_OUT": "Ra về", "PENDING": "Chờ duyệt",
                  "APPROVED": "Đã duyệt", "REJECTED": "Từ chối", "CANCELLED": "Đã hủy",
                  "ACTIVE": "Đang làm việc", "INACTIVE": "Ngừng hoạt động",
                  "LEAVE": "Nghỉ phép", "ATTENDANCE_FIX": "Giải trình công"}
    if field in {"event_type", "status", "employment_status", "request_type"}:
        return safe_xlsx_text(vocabulary.get(str(value), value))
    return safe_xlsx_text(value)


def build_xlsx_report(report: dict, created_by: str = "", *, use_template: bool = True) -> bytes:
    """Render a fixed, macro-free report workbook from an already scoped report."""
    kind = report["kind"]
    report_fields = {
        "ATTENDANCE": ["employee_code", "full_name", "department_id", "recorded_days", "present_days", "incomplete_days", "absent_days", "worked_minutes"],
        "LEAVE": ["employee_code", "full_name", "department_id", "approved_days", "pending_requests", "rejected_requests", "remaining_annual_days"],
        "ATTENDANCE_FIX": ["attendance_fix_id", "employee_code", "full_name", "department_id", "work_date", "event_type", "requested_at", "status", "reason"],
        "APPROVAL_QUEUE": ["request_type", "request_id", "employee_code", "full_name", "department_id", "request_date", "detail", "created_at"],
        "HEADCOUNT": ["department_id", "department_name", "employment_status", "employee_count"],
        "MY_PAYSLIP": ["employee_code", "full_name", "period_year", "period_month", "currency_code", "base_salary", "standard_work_days", "actual_work_days", "allowance_amount", "overtime_amount", "deduction_amount", "gross_salary", "insurance_deduction", "taxable_income", "personal_income_tax", "other_deductions", "net_salary"],
        "PAYROLL_SUMMARY": ["department_id", "department_name", "currency_code", "employee_count", "base_salary", "allowance_amount", "overtime_amount", "deduction_amount", "gross_salary", "insurance_deduction", "taxable_income", "personal_income_tax", "other_deductions", "net_salary"],
    }
    columns = report_fields[kind]
    metric_fields = [field for field in columns if field.endswith("_days") or field.endswith("_requests") or field in {"recorded_days", "present_days", "incomplete_days", "absent_days", "worked_minutes", "employee_count", "base_salary", "allowance_amount", "overtime_amount", "deduction_amount", "gross_salary", "insurance_deduction", "taxable_income", "personal_income_tax", "other_deductions", "net_salary"}]
    labels = {
        "employee_code": "Mã nhân viên", "full_name": "Họ tên", "department_id": "Mã phòng ban",
        "recorded_days": "Ngày có bản ghi", "present_days": "Đủ công", "incomplete_days": "Thiếu lượt",
        "absent_days": "Vắng", "worked_minutes": "Phút làm việc", "approved_days": "Ngày phép đã duyệt",
        "pending_requests": "Đơn chờ duyệt", "rejected_requests": "Đơn bị từ chối",
        "remaining_annual_days": "Số dư phép năm hiện tại",
        "attendance_fix_id": "Mã giải trình", "work_date": "Ngày công", "event_type": "Lượt công",
        "requested_at": "Giờ đề nghị", "status": "Trạng thái", "reason": "Lý do",
        "request_type": "Loại đề nghị", "request_id": "Mã đề nghị", "request_date": "Ngày nghiệp vụ",
        "detail": "Nội dung", "created_at": "Tạo lúc", "department_name": "Tên phòng ban",
        "employment_status": "Trạng thái nhân sự", "employee_count": "Số nhân viên",
        "period_year": "Năm lương", "period_month": "Tháng lương", "currency_code": "Mã tiền tệ",
        "base_salary": "Lương cơ bản", "standard_work_days": "Ngày công chuẩn", "actual_work_days": "Ngày công thực tế",
        "allowance_amount": "Phụ cấp", "overtime_amount": "Tiền làm thêm", "deduction_amount": "Khoản khấu trừ",
        "gross_salary": "Tổng lương GROSS", "insurance_deduction": "Khấu trừ bảo hiểm", "taxable_income": "Thu nhập chịu thuế",
        "personal_income_tax": "Thuế thu nhập cá nhân", "other_deductions": "Khấu trừ khác", "net_salary": "Lương thực nhận NET",
    }
    if use_template:
        workbook = load_workbook(Path(__file__).resolve().parents[1] / "report_templates" / f"{kind.lower()}_v1.xlsx")
        for sheet in workbook.worksheets:
            for merged in list(sheet.merged_cells.ranges):
                sheet.unmerge_cells(str(merged))
            sheet.delete_rows(1, sheet.max_row)
    else:
        workbook = Workbook()
    info = workbook.active
    info.title = "Thông tin"
    info.append(["BÁO CÁO NHÂN SỰ", ""])
    titles = {"ATTENDANCE": "Báo cáo chấm công", "LEAVE": "Báo cáo nghỉ phép",
              "ATTENDANCE_FIX": "Báo cáo giải trình công", "APPROVAL_QUEUE": "Danh sách đề nghị chờ duyệt",
              "HEADCOUNT": "Báo cáo nhân sự hiện tại", "MY_PAYSLIP": "Phiếu lương cá nhân",
              "PAYROLL_SUMMARY": "Bảng tổng hợp lương đã phát hành"}
    info["A1"] = titles[kind].upper()
    info.append(["Loại báo cáo", titles[kind]])
    info.append(["Từ ngày", report["start_date"]])
    info.append(["Đến ngày", report["end_date"]])
    info.append(["Phạm vi", {"SELF": "Bản thân", "DIRECT_REPORTS": "Nhân viên trực tiếp", "COMPANY": "Toàn công ty"}[report["scope"]]])
    info.append(["Người tạo", safe_xlsx_text(created_by)])
    now_utc = datetime.now(timezone.utc)
    info.append(["Tạo lúc (UTC)", now_utc.strftime("%Y-%m-%d %H:%M:%S %Z")])
    info.append(["Tạo lúc (Asia/Ho_Chi_Minh)", now_utc.astimezone(ZoneInfo("Asia/Ho_Chi_Minh")).strftime("%Y-%m-%d %H:%M:%S %Z")])
    info.append(["Template version", f"{kind.lower()}_v1"])
    info.append(["Ghi chú", safe_xlsx_text(report["note"])])
    info.append(["Số dòng dữ liệu", len(report["rows"])])

    data = workbook["Dữ liệu"] if "Dữ liệu" in workbook.sheetnames else workbook.create_sheet("Dữ liệu")
    data.append([labels[field] for field in columns])
    for row in report["rows"]:
        data.append([xlsx_value(field, row.get(field)) for field in columns])

    summary = workbook["Tổng hợp"] if "Tổng hợp" in workbook.sheetnames else workbook.create_sheet("Tổng hợp")
    summary.append(["Chỉ số", "Giá trị"])
    if kind in {"ATTENDANCE_FIX", "APPROVAL_QUEUE"}:
        summary.append(["Số yêu cầu trong kỳ", len(report["rows"])])
    else:
        summary.append(["Số nhân viên trong phạm vi", sum(int(row.get("employee_count", 1)) for row in report["rows"])])
    summary.append(["Loại báo cáo", titles[kind]])
    summary.append(["Kỳ", f"{report['start_date'].isoformat()} – {report['end_date'].isoformat()}"])
    if kind == "MY_PAYSLIP" and report["rows"]:
        person = report["rows"][0]
        summary.append(["Mã nhân viên", safe_xlsx_text(person.get("employee_code"))])
        summary.append(["Họ tên", safe_xlsx_text(person.get("full_name"))])
        summary.append(["Tiền tệ", safe_xlsx_text(person.get("currency_code"))])
    for field in metric_fields:
        if field in {"remaining_annual_days", "department_id", "request_id", "attendance_fix_id", "employee_count"}:
            continue  # Current balances are per employee, not additive.
        if kind in {"MY_PAYSLIP", "PAYROLL_SUMMARY"}:
            currencies = sorted({row["currency_code"] for row in report["rows"]})
            for currency in currencies:
                if not isinstance(currency, str) or not re.fullmatch(r"[A-Z]{3}", currency):
                    raise ValueError("PAYROLL_CURRENCY_UNAVAILABLE")
                values = [Decimal(str(row[field])) for row in report["rows"] if row.get(field) is not None and row["currency_code"] == currency]
                summary.append([f"{labels[field]} ({currency})", xlsx_decimal(sum(values, Decimal("0")))])
        else:
            values = [Decimal(str(row[field])) for row in report["rows"] if row.get(field) is not None]
            summary.append([labels[field], xlsx_decimal(sum(values, Decimal("0")))])
    summary.append(["Giới hạn dữ liệu", safe_xlsx_text(report["note"])])
    summary.append(["Độ chính xác Excel", "Số vượt 15 chữ số được lưu dưới dạng text chính xác."])

    header_fill = PatternFill("solid", fgColor="17365D")
    line = Side(style="thin", color="E2E8F0")
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape" if sheet is data else "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:1"
        sheet.print_area = sheet.dimensions
        sheet.oddHeader.left.text = titles[kind]
        sheet.oddHeader.right.text = "Nội bộ"
        sheet.oddFooter.left.text = "Group 3 HRMS • Asia/Ho_Chi_Minh"
        sheet.oddFooter.right.text = "Trang &P / &N"
        sheet.freeze_panes = "C2" if sheet is data else "A2"
        sheet.auto_filter.ref = sheet.dimensions if sheet is data else None
        sheet.row_dimensions[1].height = 34
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font(name="Calibri", size=11, color="243746")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(bottom=line)
                if cell.row > 1 and cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F3F6FA")
                if cell.data_type == "n":
                    cell.number_format = "#,##0.00" if isinstance(cell.value, (Decimal, float)) else "#,##0"
                if isinstance(cell.value, datetime):
                    cell.number_format = "dd/mm/yyyy hh:mm"
                elif isinstance(cell.value, date):
                    cell.number_format = "dd/mm/yyyy"
            if row[0].row > 1:
                sheet.row_dimensions[row[0].row].height = 30
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column in sheet.columns:
            width = min(48, max(16, max(len(str(cell.value or "")) for cell in column) + 3))
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width

    info.merge_cells("A1:B1")
    info["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    info.row_dimensions[1].height = 44
    for sheet in (info, summary):
        sheet.column_dimensions["A"].width = 38
        sheet.column_dimensions["B"].width = 72
    info.row_dimensions[10].height = 80
    summary.row_dimensions[summary.max_row - 1].height = 80
    for index, field in enumerate(columns, 1):
        if field in {"approved_days", "remaining_annual_days", "standard_work_days", "actual_work_days"}:
            for row in data.iter_rows(min_row=2, min_col=index, max_col=index):
                row[0].number_format = "#,##0.00"
        if field in {"reason", "detail", "full_name"}:
            data.column_dimensions[get_column_letter(index)].width = 36
        elif field in {"work_date", "request_date"}:
            data.column_dimensions[get_column_letter(index)].width = 18
        elif field in {"requested_at", "created_at"}:
            data.column_dimensions[get_column_letter(index)].width = 24
    if report["rows"]:
        table = Table(displayName="ReportData", ref=data.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        data.add_table(table)
    for row in summary.iter_rows(min_row=2):
        if str(row[0].value).startswith("Lương thực nhận NET"):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="E2EFDA")
                cell.font = Font(name="Calibri", size=12, bold=True, color="244B27")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class ReportService:
    @staticmethod
    async def _released_payroll_period(db, start: date, end: date) -> PayrollPeriod:
        from calendar import monthrange
        if start.day != 1 or end != start.replace(day=monthrange(start.year, start.month)[1]):
            raise HTTPException(status_code=422, detail="PAYROLL_REPORT_REQUIRES_ONE_CALENDAR_MONTH")
        period = await db.scalar(select(PayrollPeriod).where(
            PayrollPeriod.start_date == start,
            PayrollPeriod.end_date == end,
        ))
        if not period or period.status not in {"APPROVED", "CLOSED"}:
            raise HTTPException(status_code=404, detail="REPORT_UNAVAILABLE")
        return period

    @staticmethod
    async def visible_employee_ids(
        db, user, kind: str, scope: str | None, department_id: int | None,
        start: date | None = None, end: date | None = None,
    ) -> tuple[str, list[int]]:
        roles = user_roles(user)
        if kind == "MY_PAYSLIP":
            if scope not in {None, "SELF"} or department_id is not None:
                raise HTTPException(status_code=403, detail="AI_ACCESS_DENIED")
            report_scope(user, "SELF")
            if start is None or end is None:
                raise HTTPException(status_code=422, detail="PAYROLL_PERIOD_REQUIRED")
            await ReportService._released_payroll_period(db, start, end)
            return "SELF", [user.employee_id] if user.employee_id is not None else []
        if kind == "PAYROLL_SUMMARY":
            if not roles & {"HR", "ADMIN"} or scope not in {None, "COMPANY"}:
                raise HTTPException(status_code=403, detail="AI_ACCESS_DENIED")
            report_scope(user, "COMPANY")
            if start is None or end is None:
                raise HTTPException(status_code=422, detail="PAYROLL_PERIOD_REQUIRED")
            period = await ReportService._released_payroll_period(db, start, end)
            statement = select(PayrollLine.employee_id).join(Employee).where(
                PayrollLine.payroll_period_id == period.payroll_period_id
            )
            if department_id is not None:
                statement = statement.where(Employee.department_id == department_id)
            ids = list((await db.scalars(statement.order_by(PayrollLine.employee_id).limit(MAX_REPORT_ROWS + 1))).all())
            return "COMPANY", check_report_size(ids)
        if kind in {"HEADCOUNT", "APPROVAL_QUEUE"} and not roles & {"MANAGER", "HR", "ADMIN"}:
            raise HTTPException(status_code=403, detail="AI_ACCESS_DENIED")
        if kind == "APPROVAL_QUEUE" and scope not in {None, "DIRECT_REPORTS"}:
            raise HTTPException(403, "AI_ACCESS_DENIED")
        requested = "DIRECT_REPORTS" if kind == "APPROVAL_QUEUE" else scope
        resolved_scope = report_scope(user, requested)
        stmt = select(Employee.employee_id)
        scope_filter = employee_scope(user, resolved_scope)
        if scope_filter is not None:
            stmt = stmt.where(scope_filter)
        if department_id is not None:
            stmt = stmt.where(Employee.department_id == department_id)
        ids = list((await db.scalars(stmt.order_by(Employee.employee_id).limit(MAX_REPORT_ROWS + 1))).all())
        return resolved_scope, check_report_size(ids)

    @staticmethod
    async def generate(
        db, user, kind: str, start: date, end: date,
        department_id: int | None, scope: str | None = None,
        source_employee_ids: list[int] | None = None,
    ):
        roles = user_roles(user)
        if kind in {"MY_PAYSLIP", "PAYROLL_SUMMARY"}:
            if kind == "MY_PAYSLIP" and department_id is not None:
                raise HTTPException(status_code=403, detail="AI_ACCESS_DENIED")
            if kind == "MY_PAYSLIP":
                if scope not in {None, "SELF"}:
                    raise HTTPException(status_code=403, detail="AI_ACCESS_DENIED")
                resolved_scope = report_scope(user, "SELF")
            else:
                if not roles & {"HR", "ADMIN"} or scope not in {None, "COMPANY"}:
                    raise HTTPException(status_code=403, detail="AI_ACCESS_DENIED")
                resolved_scope = report_scope(user, "COMPANY")

            period = await ReportService._released_payroll_period(db, start, end)
            line_stmt = select(
                PayrollLine, Employee.employee_code, Employee.full_name, Employee.department_id, Department.department_name,
            ).join(Employee, Employee.employee_id == PayrollLine.employee_id).outerjoin(Department, Department.department_id == Employee.department_id).where(
                PayrollLine.payroll_period_id == period.payroll_period_id
            )
            if department_id is not None:
                line_stmt = line_stmt.where(Employee.department_id == department_id)
            if source_employee_ids is not None:
                line_stmt = line_stmt.where(PayrollLine.employee_id.in_(source_employee_ids))
            if kind == "MY_PAYSLIP":
                if user.employee_id is None:
                    rows = []
                else:
                    line_stmt = line_stmt.where(PayrollLine.employee_id == user.employee_id)
                    line = (await db.execute(line_stmt)).first()
                    if line is not None and not re.fullmatch(r"[A-Z]{3}", line[0].currency_code or ""):
                        raise HTTPException(409, "PAYROLL_CURRENCY_UNAVAILABLE")
                    rows = [] if line is None else [{
                        "employee_id": line[0].employee_id,
                        "employee_code": line[1],
                        "full_name": line[2],
                        "department_id": line[3],
                        "period_year": period.period_year,
                        "period_month": period.period_month,
                        "currency_code": line[0].currency_code,
                        "base_salary": line[0].base_salary,
                        "standard_work_days": line[0].standard_work_days,
                        "actual_work_days": line[0].actual_work_days,
                        "allowance_amount": line[0].allowance_amount,
                        "overtime_amount": line[0].overtime_amount,
                        "deduction_amount": line[0].deduction_amount,
                        "gross_salary": line[0].gross_salary,
                        "insurance_deduction": line[0].insurance_deduction,
                        "taxable_income": line[0].taxable_income,
                        "personal_income_tax": line[0].personal_income_tax,
                        "other_deductions": line[0].other_deductions,
                        "net_salary": line[0].net_salary,
                    }]
                note = "Bảng lương cá nhân lấy số liệu và mã tiền tệ đã lưu trong snapshot kỳ APPROVED/CLOSED; không quy đổi tiền hoặc tính lại lương."
            else:
                lines = check_report_size((await db.execute(line_stmt.limit(MAX_REPORT_ROWS + 1))).all())
                amount_fields = (
                    "base_salary", "allowance_amount", "overtime_amount", "deduction_amount", "gross_salary",
                    "insurance_deduction", "taxable_income", "personal_income_tax", "other_deductions", "net_salary",
                )
                groups = {}
                for line in lines:
                    currency = line[0].currency_code
                    if not re.fullmatch(r"[A-Z]{3}", currency or ""):
                        raise HTTPException(409, "PAYROLL_CURRENCY_UNAVAILABLE")
                    key = (line[3], line[4], currency)
                    group = groups.setdefault(key, {"department_id": line[3], "department_name": line[4], "currency_code": currency, "employee_count": 0, **{field: Decimal("0.00") for field in amount_fields}})
                    group["employee_count"] += 1
                    for field in amount_fields:
                        group[field] += getattr(line[0], field) or Decimal("0.00")
                rows = list(groups.values())
                note = "Tổng hợp theo phòng ban hiện tại và từng mã tiền tệ từ snapshot APPROVED/CLOSED; không gộp khác tiền tệ, quy đổi, tính lại lương hoặc suy ra lịch sử phòng ban."
            for row in rows:
                for key, value in row.items():
                    if isinstance(value, Decimal):
                        row[key] = str(value)
            return {
                "kind": kind, "scope": resolved_scope, "start_date": start, "end_date": end,
                "rows": rows, "note": note,
            }
        if kind == "HEADCOUNT" and not roles & {"MANAGER", "HR", "ADMIN"}:
            raise HTTPException(status_code=403, detail="AI_ACCESS_DENIED")
        if kind == "APPROVAL_QUEUE" and not roles & {"MANAGER", "HR", "ADMIN"}:
            raise HTTPException(status_code=403, detail="AI_ACCESS_DENIED")
        if kind == "APPROVAL_QUEUE" and scope not in {None, "DIRECT_REPORTS"}:
            raise HTTPException(403, "AI_ACCESS_DENIED")
        requested = "DIRECT_REPORTS" if kind == "APPROVAL_QUEUE" else scope
        resolved_scope = report_scope(user, requested)
        employees = select(Employee.employee_id, Employee.employee_code, Employee.full_name, Employee.department_id)
        scope_filter = employee_scope(user, resolved_scope)
        if source_employee_ids is not None:
            frozen_filter = Employee.employee_id.in_(source_employee_ids)
            scope_filter = and_(scope_filter, frozen_filter) if scope_filter is not None else frozen_filter
        if scope_filter is not None:
            employees = employees.where(scope_filter)
        if department_id is not None:
            employees = employees.where(Employee.department_id == department_id)
        visible = employees.subquery()
        if kind == "ATTENDANCE":
            summary = select(
                AttendanceDay.employee_id,
                func.count().label("recorded_days"),
                func.sum(case((AttendanceDay.attendance_status == "PRESENT", 1), else_=0)).label("present_days"),
                func.sum(case((AttendanceDay.attendance_status == "INCOMPLETE", 1), else_=0)).label("incomplete_days"),
                func.sum(case((AttendanceDay.attendance_status == "ABSENT", 1), else_=0)).label("absent_days"),
                func.sum(AttendanceDay.worked_minutes).label("worked_minutes"),
            ).join(visible, visible.c.employee_id == AttendanceDay.employee_id).where(
                AttendanceDay.work_date.between(start, end)
            ).group_by(AttendanceDay.employee_id).subquery()
            metrics = ["recorded_days", "present_days", "incomplete_days", "absent_days", "worked_minutes"]
            stmt = select(visible, *[func.coalesce(summary.c[name], 0).label(name) for name in metrics]).outerjoin(summary, summary.c.employee_id == visible.c.employee_id)
        elif kind == "LEAVE":
            summary = select(
                LeaveRequest.employee_id,
                func.sum(case((LeaveRequest.status == "APPROVED", LeaveRequest.leave_days), else_=Decimal("0"))).label("approved_days"),
                func.sum(case((LeaveRequest.status == "PENDING", 1), else_=0)).label("pending_requests"),
                func.sum(case((LeaveRequest.status == "REJECTED", 1), else_=0)).label("rejected_requests"),
            ).join(visible, visible.c.employee_id == LeaveRequest.employee_id).where(
                LeaveRequest.leave_date.between(start, end)
            ).group_by(LeaveRequest.employee_id).subquery()
            metrics = ["approved_days", "pending_requests", "rejected_requests"]
            stmt = select(visible, *[func.coalesce(summary.c[name], 0).label(name) for name in metrics],
                EmployeeLeaveBalance.remaining_days.label("remaining_annual_days")
            ).outerjoin(summary, summary.c.employee_id == visible.c.employee_id).outerjoin(
                EmployeeLeaveBalance, (EmployeeLeaveBalance.employee_id == visible.c.employee_id) & (EmployeeLeaveBalance.year == start.year)
            )
        elif kind == "ATTENDANCE_FIX":
            stmt = select(
                AttendanceFix.attendance_fix_id, visible.c.employee_id, visible.c.employee_code,
                visible.c.full_name, visible.c.department_id, AttendanceFix.work_date,
                AttendanceFix.event_type, AttendanceFix.requested_at, AttendanceFix.status,
                AttendanceFix.reason, AttendanceFix.created_at,
            ).join(visible, visible.c.employee_id == AttendanceFix.employee_id).where(
                AttendanceFix.work_date.between(start, end)
            ).order_by(AttendanceFix.work_date, visible.c.employee_code, AttendanceFix.attendance_fix_id)
        elif kind == "APPROVAL_QUEUE":
            leave_stmt = select(
                LeaveRequest.leave_request_id.label("request_id"), visible.c.employee_id,
                visible.c.employee_code, visible.c.full_name, visible.c.department_id,
                LeaveRequest.leave_date.label("request_date"), LeaveRequest.session.label("detail"),
                LeaveRequest.created_at,
            ).join(visible, visible.c.employee_id == LeaveRequest.employee_id).where(
                LeaveRequest.status == "PENDING", LeaveRequest.leave_date.between(start, end)
            )
            fix_stmt = select(
                AttendanceFix.attendance_fix_id.label("request_id"), visible.c.employee_id,
                visible.c.employee_code, visible.c.full_name, visible.c.department_id,
                AttendanceFix.work_date.label("request_date"), AttendanceFix.event_type.label("detail"),
                AttendanceFix.created_at,
            ).join(visible, visible.c.employee_id == AttendanceFix.employee_id).where(
                AttendanceFix.status == "PENDING", AttendanceFix.work_date.between(start, end)
            )
            leave_rows = [dict(row, request_type="LEAVE") for row in (await db.execute(leave_stmt.limit(MAX_REPORT_ROWS + 1))).mappings().all()]
            fix_rows = [dict(row, request_type="ATTENDANCE_FIX") for row in (await db.execute(fix_stmt.limit(MAX_REPORT_ROWS + 1))).mappings().all()]
            rows = sorted(leave_rows + fix_rows, key=lambda row: (row["request_date"], row["created_at"], row["request_type"]))
            stmt = None
        elif kind == "HEADCOUNT":
            stmt = select(
                Employee.department_id, Department.department_name, Employee.employment_status,
                func.count(Employee.employee_id).label("employee_count"),
            ).join(Department, Department.department_id == Employee.department_id).group_by(
                Employee.department_id, Department.department_name, Employee.employment_status
            ).order_by(Department.department_name, Employee.employment_status)
            if scope_filter is not None:
                stmt = stmt.where(scope_filter)
            if department_id is not None:
                stmt = stmt.where(Employee.department_id == department_id)
        else:
            raise HTTPException(status_code=422, detail="REPORT_KIND_UNSUPPORTED")
        if kind != "APPROVAL_QUEUE":
            rows = [dict(row) for row in (await db.execute(stmt.limit(MAX_REPORT_ROWS + 1))).mappings().all()]
        check_report_size(rows)
        # Decimal values remain exact in JSON and CSV.
        for row in rows:
            for key, value in row.items():
                if isinstance(value, Decimal):
                    row[key] = str(value)
        notes = {
            "ATTENDANCE": "Chỉ tổng hợp ngày công đã lưu; không suy diễn ngày chưa có bản ghi.",
            "LEAVE": "Phép được duyệt gồm mọi loại phép. Số dư là số dư hiện tại của năm, không phải số dư tại cuối kỳ.",
            "ATTENDANCE_FIX": "Chỉ gồm giải trình đã được lưu trong kỳ; lý do do người lao động cung cấp.",
            "APPROVAL_QUEUE": "Chỉ gồm đề nghị PENDING của cấp dưới trực tiếp; quyền duyệt được kiểm tra lại khi xử lý.",
            "HEADCOUNT": "Ảnh chụp trạng thái nhân sự hiện tại; không suy ra biến động lịch sử.",
        }
        return {"kind": kind, "scope": resolved_scope, "start_date": start, "end_date": end, "balance_year": start.year if kind == "LEAVE" else None, "rows": rows, "note": notes[kind]}
